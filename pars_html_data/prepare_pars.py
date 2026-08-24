import os

from openpyxl import Workbook, load_workbook

from pars_html_data.utils_pars import (
    create_xml,
    get_all_parent_tag,
    get_href_from_a,
    get_parent_tag,
    get_tag,
    get_tag_a,
    get_title_from_a,
)


async def get_pagination_max_counter(html):
    soup = await create_xml(html)
    parent = await get_parent_tag(soup, "div", "paginationWrapper")
    get_input_tag = await get_tag(parent, "input", "form-control")
    max_counter = get_input_tag.get("max")

    return max_counter


def save_to_excel(data, filename="rostender_results.xlsx"):
    """Функция для сохранения данных в Excel"""
    try:
        if os.path.exists(filename):
            wb = load_workbook(filename)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Ростендер"

            headers = ["Название", "Ссылка", "Стоимость", "Дата"]
            ws.append(headers)

        for item in data:
            ws.append([item["title"], item["href"], item["cost"], item["day"]])

        wb.save(filename)
        print(f"Добавлено {len(data)} записей в {filename}")

    except Exception as e:
        print(f"Ошибка при сохранении: {e}")


async def prepare_pars_rostender(html):
    data = []

    soup = await create_xml(html)
    parents = await get_all_parent_tag(soup, "article")
    for parent in parents:
        a = await get_tag_a(parent)
        title = await get_title_from_a(a)
        href = await get_href_from_a(a)
        div = await get_tag(
            parent, "div", "starting-price__price starting-price--price"
        )
        span = await get_tag(parent, "span", "black")

        data.append(
            {
                "title": title,
                "href": href,
                "cost": div.get_text(),
                "day": span.get_text(),
            }
        )

    return data


async def prepare_pars_b2b(html):
    soup = await create_xml(html)
    parents = await get_parent_tag(soup, "tr")
    for parent in parents:
        a = await get_tag_a(parent)
        href = await get_href_from_a(a)
        title = await get_tag(parent, "div", "search-results-title-desc")
        day = await get_tag(parent, "td", "nowrap")

        title_text = title.get_text(strip=True) if title else "N/A"
        day_text = day.get_text(strip=True) if day else "N/A"

        print(f"title={title_text} \n href = {href} \n cost= - \n day={day_text}")


async def prepare_pars_bidzaar(html):
    soup = await create_xml(html)
    parents = await get_parent_tag(soup, "li")
    for parent in parents:
        a = await get_tag_a(parent)
        href = await get_href_from_a(a)
        title = await get_tag(parent, "span", "name-item ui-name with-dot")
        day = await get_tag(parent, "div", "date")

        title_text = title.get_text(strip=True) if title else "N/A"
        day_text = day.get_text(strip=True) if day else "N/A"

        print(f"title={title_text} \n href = {href} \n cost= - \n day={day_text}")


async def prepare_pars_rostorg(html):
    soup = await create_xml(html)
    parents = await get_parent_tag(soup, "div", "search-results__item autoload-post")
    for parent in parents:
        a = await get_tag_a(parent)
        href = await get_href_from_a(a)
        title = await get_tag(
            parent, "a", "search-results__link search-results__link--description"
        )
        cost = await get_tag(parent, "p", "tablet")
        day = await get_tag(parent, "time", "search-results__time")

        title_text = title.get_text(strip=True) if title else "N/A"
        day_text = day.get_text(strip=True) if day else "N/A"
        cost_text = cost.get_text(strip=True) if cost else "N/A"

        print(
            f"title={title_text} \n href = {href} \n cost={cost_text} \n day={day_text}"
        )
